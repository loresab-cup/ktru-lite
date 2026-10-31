"""Case 5 acceptance through the running HTTP API; --live explicitly enables paid LLM runs."""

import argparse
import csv
import io
import json
import math
import os
import secrets
import sys
import time
from pathlib import Path

import httpx
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import text

import auth
from database import SessionLocal


def check_confidence(items):
    """Every algorithm proposal needs a finite percentage; human decisions are not model scores."""
    proposals = [
        i
        for i in items
        if i["status"] in ("auto", "needs_review")
        and not (i.get("attributes") or {}).get("remembered")
    ]
    valid = sum(
        isinstance(i["confidence"], (int, float))
        and not isinstance(i["confidence"], bool)
        and math.isfinite(i["confidence"])
        and 0 <= i["confidence"] <= 100
        for i in proposals
    )
    return {
        "automatic_proposals": len(proposals),
        "remembered_human_decisions": sum(
            bool((i.get("attributes") or {}).get("remembered")) for i in items
        ),
        "scored": valid,
        "unscored_or_invalid": len(proposals) - valid,
        "passed": bool(items) and valid == len(proposals),
    }


def pdf(pages):
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for _ in range(pages):
        page = writer.add_blank_page(width=595, height=842)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 40 780 Td (Notebook 16 GB) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def run(client, name, data, timeout=600):
    started = time.monotonic()
    response = client.post("/api/jobs", files={"file": (name, data)})
    response.raise_for_status()
    job_id = response.json()["job_id"]
    upload = time.monotonic() - started
    while time.monotonic() - started < timeout:
        response = client.get("/api/jobs/" + job_id)
        response.raise_for_status()
        job = response.json()
        if job["status"] in ("done", "failed", "cancelled"):
            break
        time.sleep(0.5)
    else:
        client.post("/api/jobs/" + job_id + "/cancel").raise_for_status()
        raise TimeoutError("Acceptance job exceeded deadline")
    elapsed = time.monotonic() - started
    items = []
    for offset in range(0, job["total"], 200):
        page = client.get("/api/jobs/" + job_id + "/items", params={"limit": 200, "offset": offset})
        page.raise_for_status()
        items += page.json()["items"]
    return (
        job,
        items,
        {
            "job_id": job_id,
            "status": job["status"],
            "rows": len(items),
            "upload_seconds": round(upload, 3),
            "seconds": round(elapsed, 3),
            "fallback": job.get("fallback", 0),
            "words_only": job.get("words_only", 0),
        },
    )


def manual_correction(client, items):
    """Another existing code (the second candidate) goes in by hand and is still there when the
    position is read again."""
    item = next(i for i in items if len(i.get("candidates") or []) > 1)
    other = next(c["code"] for c in item["candidates"] if c["code"] != item["ktru_code"])
    changed = client.patch("/api/items/" + str(item["id"]), json={"ktru_code": other})
    changed.raise_for_status()
    again = client.get("/api/items/" + str(item["id"])).json()
    return {
        "row_number": item["row_number"],
        "from": item["ktru_code"],
        "to": other,
        "passed": changed.json()["status"] == "corrected"
        and again["ktru_code"] == other
        and again["status"] == "corrected",
    }


def export_table(fmt, content):
    if fmt == "csv":
        return list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(content), read_only=True).active
    return [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]


def check_export(client, job_id, fmt, corrected):
    """The file has every row and the code the person set, not just HTTP 200."""
    response = client.get("/api/jobs/" + job_id + "/export", params={"format": fmt})
    response.raise_for_status()
    table = export_table(fmt, response.content)
    header, rows = table[0], table[1:]
    line, code = header.index("Строка"), header.index("Код КТРУ")
    by_row = {str(r[line]): r[code] for r in rows}
    return {
        "rows": len(rows),
        "corrected_code_in_file": by_row.get(str(corrected["row_number"])) == corrected["to"],
        "passed": len(rows) == 1000 and by_row.get(str(corrected["row_number"])) == corrected["to"],
    }


def check_isolation(api, job_id):
    """Another account sees neither the job nor its export."""
    login = "accept-other-" + secrets.token_hex(4)
    with SessionLocal() as s:
        token = auth.start_session(s, auth.create_user(s, login, secrets.token_urlsafe(24)), 1)
    try:
        with httpx.Client(base_url=api, timeout=60) as other:
            other.cookies.set("ktru_session", token)
            answers = {
                "job": other.get("/api/jobs/" + job_id).status_code,
                "export": other.get("/api/jobs/" + job_id + "/export").status_code,
                "listed": job_id in {j["id"] for j in other.get("/api/jobs").json()},
            }
    finally:
        with SessionLocal() as s:
            auth.remove_user(s, login)
    return answers | {
        "passed": answers["job"] in (403, 404)
        and answers["export"] in (403, 404)
        and not answers["listed"]
    }


def main():
    args = argparse.ArgumentParser(description=__doc__)
    args.add_argument("--live", action="store_true")
    options = args.parse_args()
    report = {"case": 5, "live": options.live, "checks": {}, "limitations": []}
    login = "accept-" + secrets.token_hex(6)
    with SessionLocal() as s:
        user = auth.create_user(s, login, secrets.token_urlsafe(24))
        token = auth.start_session(s, user, 1)
    try:
        with httpx.Client(
            base_url=os.getenv("ACCEPTANCE_API", "http://backend:8000"), timeout=60
        ) as client:
            client.cookies.set("ktru_session", token)
            import privacy

            for purpose in ("account", "files"):
                client.post(
                    "/api/auth/consents",
                    json={
                        "purpose": purpose,
                        "accepted": True,
                        "version": privacy.VERSION,
                    },
                ).raise_for_status()
            health = client.get("/api/health").json()
            from tasks import app
            from database import engine

            workers = bool(app.control.ping(timeout=5))
            with engine.connect() as connection:
                extension = connection.scalar(
                    text("SELECT extversion FROM pg_extension WHERE extname='vector'")
                )
            report["checks"]["stack"] = {
                "passed": health.get("catalog") == "ready"
                and health.get("queue") == "celery"
                and health.get("vector_search") == "pgvector"
                and workers
                and bool(extension),
                "catalog_codes": health.get("codes"),
                "worker_available": workers,
                "pgvector_version": extension,
            }
            if not options.live:
                report["limitations"].append(
                    "Live classification, attributes and PDF processing not run; use --live."
                )
            else:
                print(
                    "Running PDF 51 pages, CSV/XLSX, attributes and 1000-row classification",
                    file=sys.stderr,
                    flush=True,
                )
                _, _, result = run(client, "acceptance-51-pages.pdf", pdf(51))
                report["checks"]["pdf_51_pages"] = result | {
                    "passed": result["status"] == "done"
                    and result["rows"] == 51
                    and result["upload_seconds"] < 30
                }
                brands = [
                    ("Ноутбук", "Lenovo", "ThinkPad T14"),
                    ("Принтер", "Canon", "i-SENSYS LBP6030B"),
                    ("Монитор", "Samsung", "S24R350"),
                    ("Мышь", "Logitech", "M185"),
                    ("Клавиатура", "Logitech", "K120"),
                    ("Смартфон", "Samsung", "Galaxy A54"),
                    ("Телевизор", "LG", "43UR78006LK"),
                    ("Маршрутизатор", "TP-Link", "Archer C6"),
                    ("Проектор", "Epson", "EB-X49"),
                    ("Планшет", "Apple", "iPad Air"),
                    ("Сканер", "Canon", "CanoScan LiDE 400"),
                    ("Наушники", "Sony", "WH-1000XM4"),
                ]
                buffer = io.StringIO()
                writer = csv.writer(buffer, delimiter=";")
                writer.writerow(["Наименование"])
                for name, brand, model in brands:
                    writer.writerow([f"{name}. Бренд: {brand}. Модель: {model}"])
                _, items, result = run(
                    client, "acceptance-attributes.csv", buffer.getvalue().encode("utf-8-sig")
                )
                hits = dict.fromkeys(("name", "brand", "model"), 0)
                for row, expected in zip(items, brands, strict=True):
                    for key, value in zip(hits, expected, strict=True):
                        normalize = lambda v: " ".join(v.lower().replace("ё", "е").split())
                        hits[key] += normalize(row["attributes"].get(key, "") or "") == normalize(
                            value
                        )
                report["checks"]["attributes"] = result | {
                    "hits": hits,
                    "denominator": 12,
                    "passed": all(h >= 11 for h in hits.values()) and result["fallback"] == 0,
                }
                root = Path(os.getenv("ACCEPTANCE_DATA", "/acceptance-data")) / "organizers"
                source = root / "Спецификация_1000_без_КТРУ.xlsx"
                reference = {
                    r["row_number"]: r
                    for r in map(
                        json.loads,
                        (root / "Спецификация_1000_reference.jsonl")
                        .read_text(encoding="utf-8")
                        .splitlines(),
                    )
                }
                _, items, result = run(client, "acceptance-1000.xlsx", source.read_bytes())
                exact = sum(
                    i["ktru_code"] == reference[i["row_number"]]["ktru_code"] for i in items
                )
                report["checks"]["classification"] = result | {
                    "exact": exact,
                    "passed": result["status"] == "done"
                    and len(items) == 1000
                    and exact >= 850
                    and result["fallback"] == 0
                    and result["words_only"] == 0,
                }
                report["checks"]["performance"] = {
                    "seconds": result["seconds"],
                    "limit": 300,
                    "passed": result["seconds"] < 300,
                }
                report["checks"]["confidence"] = check_confidence(items)
                corrected = manual_correction(client, items)
                report["checks"]["manual_correction"] = corrected
                for fmt in ["csv", "xlsx"]:
                    report["checks"]["export_" + fmt] = check_export(
                        client, result["job_id"], fmt, corrected
                    )
                report["checks"]["isolation"] = check_isolation(
                    os.getenv("ACCEPTANCE_API", "http://backend:8000"), result["job_id"]
                )
                report["limitations"].append(
                    "Attribute set is 12 handcrafted examples; not an independent organizer benchmark. PDF is text, not a scanned document."
                )
        report["passed"] = options.live and all(c["passed"] for c in report["checks"].values())
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if report["passed"] else 1
    except Exception as exc:  # noqa: BLE001 - CLI boundary must report every failed step as JSON.
        report["passed"] = False
        report["error"] = f"{type(exc).__name__}: {exc}"
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 1
    finally:
        with SessionLocal() as s:
            auth.remove_user(s, login)


if __name__ == "__main__":
    raise SystemExit(main())
