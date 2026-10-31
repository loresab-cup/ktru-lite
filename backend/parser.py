"""Specification file -> list of positions (one Item per product row).

Formats: xlsx (every sheet), csv (utf-8 / cp1251, ; , or tab), docx (tables, then paragraphs),
pdf (tables of up to 200 pages, or its text lines when it has no tables), txt (one position
per line).
"""

import csv
import io
import logging
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ITEMS = 5000
MAX_CELLS = 300_000
MAX_PDF_PAGES = 200
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024  # zip bomb guard for xlsx and docx
MAX_CELL_CHARS = 2000
# Control characters a file may carry: PostgreSQL refuses NUL in text, Excel the rest
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

# Extension -> first bytes of a valid file (None = plain text, no signature)
SIGNATURES = {".xlsx": b"PK", ".docx": b"PK", ".pdf": b"%PDF", ".csv": None, ".txt": None}

HEADER_WORDS = ("наименование", "товар", "характеристик", "кол-во", "количество")
NAME_WORDS = ("наименование", "товар", "название", "предмет", "позиция")
NOT_NAME_PREFIXES = ("код", "№", "номер", "артикул", "окпд", "ктру", "товарный знак")
# Columns whose values describe the product and go into the position text with the name
DETAIL_WORDS = ("характерист", "бренд", "марка", "модель", "производител", "описание")
SKIP_PREFIXES = ("итого", "всего", "итог")
MAX_HEADER_CELL = 40
# Columns with instructions for the bidder, not product data (EIS forms)
SERVICE_WORDS = (
    "используется",
    "применение характеристик",
    "инструкция",
    "обоснование",
    "национальный режим",
)
# EIS «Описание объекта закупки»: a product takes several rows, one characteristic per row;
# the rows after the first have empty «№» and name cells
EIS_CHARACTERISTIC = "наименование характеристики"
EIS_VALUE = "значени"  # «Текстовое описание значения», «Конкретное значение»
EIS_CHARACTERISTIC_UNIT = "единица измерения характеристики"
EIS_ITEM_UNIT = ("ед. изм", "единица измерения товара")
EIS_QUANTITY = ("кол-во", "количество")
# The new form writes a range in four columns: «>, >=» «Min» «<, <=» «Max»
EIS_BOUNDS = ((">, >=", "min_sign"), ("Min", "min"), ("<, <=", "max_sign"), ("Max", "max"))
EIS_SIGNS = {">=": "≥", ">": ">", "<=": "≤", "<": "<"}
NUMBERING = re.compile(r"^\s*\d{1,4}\s*[.)]\s+")

logging.getLogger("pypdf").setLevel(logging.ERROR)  # broken fonts spam warnings
logging.getLogger("pdfminer").setLevel(logging.ERROR)


class ParseError(ValueError):
    """A file the user must fix; the message is shown as is (Russian)."""


@dataclass
class Item:
    row_number: int
    text: str
    columns: dict[str, str] = field(default_factory=dict)


def check_signature(ext: str, head: bytes) -> None:
    if ext not in SIGNATURES:
        raise ParseError("Формат не поддерживается: загрузите xlsx, csv, docx, pdf или txt")
    signature = SIGNATURES[ext]
    if signature is not None and not head.startswith(signature):
        raise ParseError(f"Файл не похож на {ext}: содержимое не совпадает с расширением")


def parse_file(path: Path) -> list[Item]:
    ext = path.suffix.lower()
    data = path.read_bytes()
    if len(data) > MAX_FILE_BYTES:
        raise ParseError("Файл больше 10 МБ")
    check_signature(ext, data[:8])
    readers = {".xlsx": read_xlsx, ".csv": read_csv, ".docx": read_docx, ".pdf": read_pdf}
    items = readers[ext](path) if ext in readers else parse_lines(decode_text(data).splitlines())
    if not items:
        raise ParseError("Не найдено ни одной позиции: проверьте колонку с наименованием")
    if len(items) > MAX_ITEMS:
        raise ParseError(f"В файле {len(items)} позиций, допустимо не больше {MAX_ITEMS}")
    return items


# --- readers -------------------------------------------------------------------------------


def decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


def check_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ParseError("Файл повреждён: откройте и сохраните его заново") from exc
    if total > MAX_UNZIPPED_BYTES:
        raise ParseError("Файл слишком большой после распаковки")


def read_xlsx(path: Path) -> list[Item]:
    from openpyxl import load_workbook

    check_zip(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # broken zip or xml inside
        raise ParseError("Файл xlsx повреждён: откройте и сохраните его заново") from exc
    items: list[Item] = []
    cells = 0
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_CELLS:
                    raise ParseError("В файле слишком много ячеек")
                rows.append(row)
            items.extend(parse_rows(rows))
    finally:
        workbook.close()
    return items


def read_csv(path: Path) -> list[Item]:
    text = decode_text(path.read_bytes())
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";"
    try:
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error as exc:  # a NUL byte, a cell over 128 KB: not a table to read
        raise ParseError("Файл csv повреждён: сохраните его заново как CSV (UTF-8)") from exc
    if sum(len(row) for row in rows) > MAX_CELLS:
        raise ParseError("В файле слишком много ячеек")
    return parse_rows(rows)


def read_docx(path: Path) -> list[Item]:
    from docx import Document

    check_zip(path)
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ParseError("Файл docx повреждён: откройте и сохраните его заново") from exc
    items: list[Item] = []
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        items.extend(parse_rows(rows))
    if items:
        return items
    return parse_lines([paragraph.text for paragraph in document.paragraphs])


def read_pdf(path: Path) -> list[Item]:
    """A specification PDF is usually a table: a name of two lines in one cell is one position,
    not two (a 59-page table of 291 rows gave 1344 «positions» line by line). Without tables
    every text line is a position, as before; the lines come from pypdf, which reads fonts that
    pdfminer (inside pdfplumber) sometimes cannot."""
    import pdfplumber
    from pypdf import PdfReader

    from pdf_tables import extract_tables

    try:
        reader = PdfReader(str(path))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ParseError(f"В PDF больше {MAX_PDF_PAGES} страниц")
        rows, pending = [], False
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                for number, (table, open_end) in enumerate(extract_tables(page)):
                    page_rows = [[cell or "" for cell in row] for row in table if any(row)]
                    if rows and page_rows:
                        known_headers = {
                            tuple(clean(c) for c in row)
                            for row in rows
                            if any(is_header_cell(clean(c)) for c in row)
                        }
                        while page_rows and tuple(clean(c) for c in page_rows[0]) in known_headers:
                            page_rows.pop(0)
                    continuation = page_rows and (
                        continues(rows, page_rows[0])
                        or (
                            pending
                            and rows
                            and len(page_rows[0]) == len(rows[-1])
                            and any(not c for c in page_rows[0])
                        )
                    )
                    if number == 0 and continuation:
                        rows[-1] = [f"{a} {b}".strip() for a, b in zip(rows[-1], page_rows.pop(0))]
                    rows += page_rows
                    pending = open_end
                page.close()  # frees the parsed page: 200 pages stay small in memory
        if not rows:
            lines = [
                line for page in reader.pages for line in (page.extract_text() or "").split("\n")
            ]
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("PDF повреждён или защищён паролем") from exc
    if sum(len(row) for row in rows) > MAX_CELLS:
        raise ParseError("Слишком большая таблица в PDF")
    if rows:
        return parse_rows(without_repeated_headers(rows))
    return parse_lines(lines)


def continues(rows: list[list[str]], first: list[str]) -> bool:
    """Word splits a long row over a page break: its rest opens the next page without the number
    of the first column. Such a row belongs to the last row of the previous page."""
    return (
        bool(rows)
        and len(first) == len(rows[-1])
        and len(first) > 1
        and not first[0].strip()
        and bool(rows[-1][0].strip())
    )


def without_repeated_headers(rows: list[list[str]]) -> list[list[str]]:
    """A table that goes over many pages repeats its header on each of them."""
    seen, kept = set(), []
    for row in rows:
        key = tuple(clean(cell) for cell in row)
        if key in seen and any(is_header_cell(cell) for cell in key):
            continue
        seen.add(key)
        kept.append(row)
    return kept


# --- rows and lines -> items ---------------------------------------------------------------


def clean(value) -> str:
    text = CONTROL.sub(" ", "" if value is None else str(value))
    return " ".join(text.split())[:MAX_CELL_CHARS]


def is_skipped(text: str) -> bool:
    lowered = text.lower()
    return not text or lowered.startswith(SKIP_PREFIXES) or text.replace(" ", "").isdigit()


def is_header_cell(cell: str) -> bool:
    lowered = cell.lower()
    return 0 < len(lowered) <= MAX_HEADER_CELL and any(w in lowered for w in HEADER_WORDS)


def is_name_cell(cell: str) -> bool:
    lowered = cell.lower()
    return (
        0 < len(lowered) <= MAX_HEADER_CELL
        and not lowered.startswith(NOT_NAME_PREFIXES)
        and any(word in lowered for word in NAME_WORDS)
    )


SINGLE_HEADERS = {"наименование", "наименование товара", "наименование позиции", "название товара"}


def find_header(rows: list[list[str]]) -> int | None:
    for index, row in enumerate(rows[:30]):
        filled = [cell for cell in row if cell]
        # A one-column list titled «Наименование» has a header too: it is not a product
        single_name = len(filled) == 1 and filled[0].lower() in SINGLE_HEADERS
        if single_name or (len(filled) >= 2 and any(is_header_cell(c) for c in row)):
            return index
    return None


def name_column(header: list[str], body: list[list[str]]) -> int:
    for index, cell in enumerate(header):
        if is_name_cell(cell):
            return index
    # No name header: the column with the most text
    width = max((len(row) for row in body), default=1)
    totals = [sum(len(row[i]) for row in body if i < len(row)) for i in range(width)]
    return totals.index(max(totals))


def parse_rows(raw_rows) -> list[Item]:
    """Table rows -> items; the header row is found by words like «наименование»."""
    rows = [[clean(value) for value in row] for row in raw_rows]
    header_index = find_header(rows)
    if header_index is None:
        # No header: a table of one text column is the same as lines of text
        # Cells are joined with '. ': the name stays apart from the characteristics
        return parse_lines([". ".join(cell for cell in row if cell) for row in rows])
    header = rows[header_index]
    body = rows[header_index + 1 :]
    if any(EIS_CHARACTERISTIC in cell.lower() for cell in header):
        return parse_eis_rows(header, body, header_index + 2)
    name_index = name_column(header, body)
    details = [
        i
        for i, cell in enumerate(header)
        if i != name_index and any(w in cell.lower() for w in DETAIL_WORDS) and not is_service(cell)
    ]
    header_key = [cell.lower() for cell in header if cell]
    items = []
    for offset, row in enumerate(body):
        name = row[name_index] if name_index < len(row) else ""
        if is_skipped(name) or [cell.lower() for cell in row if cell] == header_key:
            continue
        parts = [name, *(row[i] for i in details if i < len(row) and row[i])]
        columns = {
            (title or f"Колонка {i + 1}"): (row[i] if i < len(row) else "")
            for i, title in enumerate(header)
        }
        items.append(Item(header_index + offset + 2, ". ".join(parts), columns))
    return items


def is_service(title: str) -> bool:
    return any(word in title.lower() for word in SERVICE_WORDS)


def first_column(header: list[str], words: tuple[str, ...]) -> int | None:
    return next((i for i, t in enumerate(header) if any(w in t.lower() for w in words)), None)


def at(row: list[str], index: int | None) -> str:
    return row[index] if index is not None and index < len(row) else ""


def eis_unit(text: str) -> str:
    """«л; дм[3*]» -> «л»: EIS lists synonyms of one unit, the first one is enough."""
    return text.replace("[3*]", "³").replace("[2*]", "²").split(";")[0].strip()


def eis_columns(header: list[str], body: list[list[str]]) -> dict:
    # «Наименование товара, работы, услуги, код КТРУ» is longer than a usual header cell
    names = [
        i
        for i, t in enumerate(header)
        if t.lower().startswith("наименование") and "характеристик" not in t.lower()
    ]
    return {
        "number": next((i for i, t in enumerate(header) if t.startswith("№")), None),
        "name": names[0] if names else name_column(header, body),
        "characteristic": first_column(header, (EIS_CHARACTERISTIC,)),
        "values": [i for i, t in enumerate(header) if EIS_VALUE in t.lower() and not is_service(t)],
        "characteristic_unit": first_column(header, (EIS_CHARACTERISTIC_UNIT,)),
        "unit": first_column(header, EIS_ITEM_UNIT),
        "quantity": first_column(header, EIS_QUANTITY),
        **{key: header.index(title) if title in header else None for title, key in EIS_BOUNDS},
    }


def eis_range(row: list[str], columns: dict) -> str:
    """The bound columns of the new EIS form: «>=» 92 «<» 95 -> «≥ 92 и < 95»."""
    bounds = []
    for sign_key, value_key, default in (("min_sign", "min", "≥"), ("max_sign", "max", "≤")):
        value = at(row, columns[value_key])
        if value:
            bounds.append(f"{EIS_SIGNS.get(at(row, columns[sign_key]), default)} {value}")
    return " и ".join(bounds)


def add_eis_characteristic(found: list[tuple[str, list[str]]], row: list[str], columns) -> None:
    """One row of an EIS form -> («Вес», [«100 Грамм»]) in found. A row with a value but no
    characteristic name is one more value of the characteristic above («Плазма крови»)."""
    name = at(row, columns["characteristic"]).strip(" :")
    texts = [at(row, i) for i in columns["values"] if at(row, i)]
    value = " ".join([*texts, eis_range(row, columns)]).strip(" ,;")
    unit = eis_unit(at(row, columns["characteristic_unit"]))
    value = f"{value} {unit}".strip() if value else ""
    if name:
        found.append((name, [value] if value else []))
    elif value and found:
        found[-1][1].append(value)


def eis_text(found: list[tuple[str, list[str]]]) -> list[str]:
    return [f"{name}: {', '.join(values)}" if values else name for name, values in found]


def parse_eis_rows(header: list[str], body: list[list[str]], first_row: int) -> list[Item]:
    """EIS form: rows without «№» and name are more characteristics of the product above."""
    columns = eis_columns(header, body)
    header_key = [cell.lower() for cell in header if cell]
    positions: list[tuple[int, list[str], list]] = []  # row number, row, characteristics
    for offset, row in enumerate(body):
        name = at(row, columns["name"])
        if not name:
            if positions and not at(row, columns["number"]):
                add_eis_characteristic(positions[-1][2], row, columns)
            continue
        if is_skipped(name) or [cell.lower() for cell in row if cell] == header_key:
            continue
        positions.append((first_row + offset, row, []))
        add_eis_characteristic(positions[-1][2], row, columns)
    return [
        eis_item(number, row, eis_text(found), header, columns) for number, row, found in positions
    ]


def eis_item(number: int, row: list[str], found: list[str], header: list[str], columns) -> Item:
    name = at(row, columns["name"])
    unit = eis_unit(at(row, columns["unit"]))
    text = ". ".join(part for part in (name, "; ".join(found)) if part)
    if unit:
        text += f". Единица измерения: {unit}"
    kept = {
        at(header, columns["number"]): at(row, columns["number"]),
        at(header, columns["name"]): name,
        "Характеристики": "; ".join(found),
        at(header, columns["unit"]): unit,
        at(header, columns["quantity"]): at(row, columns["quantity"]),
    }
    return Item(number, text, {title: value for title, value in kept.items() if title})


def parse_lines(lines: list[str]) -> list[Item]:
    """Plain text: every non-empty line is a position; numbering «1.» is cut off."""
    items = []
    for number, line in enumerate(lines, start=1):
        text = clean(NUMBERING.sub("", line))
        lowered = text.lower()
        header_like = (
            len(text) <= MAX_HEADER_CELL * 2 and sum(word in lowered for word in HEADER_WORDS) >= 2
        )
        if is_skipped(text) or header_like or len(text) < 3:
            continue
        items.append(Item(number, text, {"Текст": text}))
    return items
