"""Regression tests for malformed providers, budgets and API responsiveness."""

import inspect
import json
import threading
from concurrent.futures import ThreadPoolExecutor

import httpx
import numpy as np
import pytest
from fastapi.testclient import TestClient

import embedder
import extractor
import parser


def test_single_column_name_header_is_not_a_product(tmp_path):
    path = tmp_path / "single.csv"
    path.write_text("Наименование\nКлавиатура\nМышь\n", encoding="utf-8-sig")
    rows = parser.parse_file(path)
    assert [(row.row_number, row.text) for row in rows] == [(2, "Клавиатура"), (3, "Мышь")]
    assert parser.parse_rows([["Мыло"], ["Бумага"]])[0].text == "Мыло"


def test_prompt_keeps_a_name_for_services_and_works():
    """YandexGPT and GigaChat gave «Услуги частной охраны» an empty name, so a right code at
    99 % never went to auto; an empty name is only for headings, totals and notes."""
    assert "Услуги и работы — тоже позиции КТРУ" in extractor.SYSTEM_PROMPT
    assert "не товар, не работа и не услуга" in extractor.SYSTEM_PROMPT


@pytest.fixture
def isolated_api(api_database):
    return api_database


@pytest.mark.parametrize("value", [float("nan"), float("inf"), 0.0])
def test_invalid_vectors_fall_back_without_poisoning_cache(monkeypatch, value):
    monkeypatch.setattr(embedder, "_cache", {})
    vector = np.full((1, embedder.DIM), value).tolist()
    # NaN/Inf are nonstandard JSON, but some gateways still serialize them.
    payload = json.dumps({"embedding": vector[0]}).encode()
    with httpx.Client(
        transport=httpx.MockTransport(lambda _: httpx.Response(200, content=payload))
    ) as client:
        monkeypatch.setattr(embedder, "_client", client)
        monkeypatch.setattr(embedder, "YANDEX_KEY", "test-key")
        monkeypatch.setattr(embedder, "YANDEX_FOLDER", "b1gfolder")
        assert embedder.embed(["Лампа"]) is None
    assert not embedder._cache


@pytest.mark.parametrize("text", ["\n=1+1", "  =1+1", "\t=1+1", "\r@SUM(1)", " +1", "-1"])
def test_export_escapes_obscured_formulas(isolated_api, text):
    assert exporting.cell(text) == "'" + text
    assert exporting.cell(123) == 123
    assert exporting.cell("Бумага") == "Бумага"


def test_blocked_upload_does_not_block_health(isolated_api, monkeypatch):
    api = isolated_api
    monkeypatch.setattr(embedder, "state", lambda: "on")
    entered, release = threading.Event(), threading.Event()
    factory = database.SessionLocal

    def slow_session():
        session = factory()
        commit = session.commit

        def slow_commit():
            entered.set()
            assert release.wait(5)
            commit()

        session.commit = slow_commit
        return session

    monkeypatch.setattr(database, "SessionLocal", slow_session)
    with TestClient(api.app) as client, ThreadPoolExecutor(2) as pool:
        upload = pool.submit(client.post, "/api/jobs", files={"file": ("a.txt", "Бумага".encode())})
        assert entered.wait(3)
        try:
            health = pool.submit(client.get, "/api/health")
            assert health.result(timeout=2).status_code == 200
        finally:
            release.set()
        assert upload.result(timeout=3).status_code == 202


def test_upload_route_uses_threadpool(isolated_api):
    assert not inspect.iscoroutinefunction(routes.jobs.create_job)


@pytest.mark.parametrize(
    "filename,content,status", [("a.txt", b"", 400), ("a.exe", b"MZ", 400), ("a.pdf", b"text", 400)]
)
def test_invalid_uploads(isolated_api, filename, content, status):
    with TestClient(isolated_api.app) as client:
        response = client.post("/api/jobs", files={"file": (filename, content)})
    assert response.status_code == status


def test_invalid_edit_payload_is_rejected(isolated_api):
    with TestClient(isolated_api.app) as client:
        assert client.put("/api/items/1/attributes", json={"name": " "}).status_code == 422
        assert client.put("/api/items/1/attributes", json={"name": "x" * 301}).status_code == 422
        assert client.get("/api/jobs", params={"limit": -1}).status_code == 422


@pytest.mark.parametrize("extension", ["txt", "csv", "xlsx", "docx"])
def test_supported_formats_without_external_fixtures(tmp_path, extension):
    path = tmp_path / f"spec.{extension}"
    if extension == "txt":
        path.write_text("Бумага А4\nЛампа 10 Вт", encoding="utf-16")
    elif extension == "csv":
        path.write_text("Наименование;Кол-во\nБумага А4;1\nЛампа 10 Вт;2", encoding="utf-8-sig")
    elif extension == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        book.active.append(["Наименование", "Кол-во"])
        book.active.append(["Бумага А4", 1])
        book.active.append(["Лампа 10 Вт", 2])
        book.save(path)
        book.close()
    else:
        from docx import Document

        document = Document()
        table = document.add_table(rows=3, cols=2)
        for row, values in zip(
            table.rows,
            [("Наименование", "Кол-во"), ("Бумага А4", "1"), ("Лампа 10 Вт", "2")],
            strict=True,
        ):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
        document.save(path)
    assert [item.text for item in parser.parse_file(path)] == ["Бумага А4", "Лампа 10 Вт"]


def test_empty_pdf_is_reported_as_no_positions(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(path)
    with pytest.raises(parser.ParseError, match="Не найдено"):
        parser.parse_file(path)


def test_1000_dataset_matches_parser_and_is_reproducible(tmp_path):
    import sys
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
    from generate_chat_data import generate

    expected = generate(tmp_path)
    content = (tmp_path / "positions_1000.jsonl").read_bytes()
    actual = parser.parse_file(tmp_path / "positions_1000.csv")
    assert len(actual) == len(expected) == 1000
    assert [vars(item) for item in actual] == expected
    generate(tmp_path)
    assert (tmp_path / "positions_1000.jsonl").read_bytes() == content


def test_oversize_upload_is_rejected(isolated_api):
    with TestClient(isolated_api.app) as client:
        response = client.post(
            "/api/jobs", files={"file": ("large.txt", b"x" * (parser.MAX_FILE_BYTES + 1))}
        )
    assert response.status_code == 413


def test_too_many_positions_is_rejected(tmp_path):
    path = tmp_path / "large.txt"
    path.write_text("Бумага А4\n" * 5001, encoding="utf-8")
    with pytest.raises(parser.ParseError, match="5000"):
        parser.parse_file(path)


def test_5000_positions_are_accepted(tmp_path):
    path = tmp_path / "large.csv"
    path.write_text(
        "Наименование;Характеристики\n" + "Бумага А4;80 г/м2\n" * 5000, encoding="utf-8"
    )
    rows = parser.parse_file(path)
    assert len(rows) == 5000
    assert rows[-1].row_number == 5001
    assert "80 г/м2" in rows[-1].text


def test_bulk_revert_accepts_5000_ids_but_rejects_5001():
    from pydantic import ValidationError

    from routes.items import RevertIn

    assert len(RevertIn(ids=list(range(1, 5001))).ids) == 5000
    with pytest.raises(ValidationError):
        RevertIn(ids=list(range(1, 5002)))


def test_job_fallback_count_handles_mixed_attribute_sources(isolated_api):
    import uuid

    api = isolated_api
    job_id = uuid.uuid4()
    with database.SessionLocal() as session:
        session.add(models.Job(id=job_id, filename="synthetic.csv"))
        session.flush()
        for number, attributes in enumerate(
            [
                {"source": "rules", "fallback": "no_key"},
                {"source": "llm"},
                {},
                {"source": "manual"},
                {"source": "manual", "fallback": "no_key"},  # rules, then fixed by a person
            ]
        ):
            session.add(
                models.Item(
                    job_id=job_id,
                    row_number=number + 1,
                    raw_text="Бумага",
                    attributes=attributes,
                    status="needs_review",
                )
            )
        session.commit()
    with TestClient(api.app) as client:
        body = client.get(f"/api/jobs/{job_id}").json()
    assert body["fallback"] == 2
    assert body["counts"]["needs_review"] == 5


class FakeCatalog:
    def get(self, code):
        return {"code": code, "name": f"Name {code}"}


def seeded_job(api, rows):
    """rows: (status, code, confidence); returns the job id and the item ids in row order."""
    import uuid

    job_id = uuid.uuid4()
    with database.SessionLocal() as session:
        session.add(models.Job(id=job_id, filename="synthetic.csv"))
        session.flush()
        items = [
            models.Item(
                job_id=job_id,
                row_number=number + 1,
                raw_text="Бумага",
                attributes={},
                status=status,
                ktru_code=code,
                ktru_name=f"Name {code}" if code else None,
                confidence=conf,
            )
            for number, (status, code, conf) in enumerate(rows)
        ]
        session.add_all(items)
        session.commit()
        return job_id, [item.id for item in items]


def test_list_is_sorted_by_code_confidence_and_status(isolated_api):
    api = isolated_api
    job_id, _ = seeded_job(
        api,
        [
            ("accepted", "B", 90.0),
            ("needs_review", None, 20.0),
            ("auto", "A", 99.0),
            ("corrected", "C", 100.0),
        ],
    )

    def codes(**params):
        body = client.get(f"/api/jobs/{job_id}/items", params=params).json()
        return [item["ktru_code"] for item in body["items"]]

    with TestClient(api.app) as client:
        assert codes() == ["B", None, "A", "C"]  # the file order
        assert codes(sort="code") == ["A", "B", "C", None]  # no code goes last
        assert codes(sort="code", order="desc") == ["C", "B", "A", None]
        assert codes(sort="confidence", order="desc") == ["C", "A", "B", None]
        assert codes(sort="status") == [None, "A", "C", "B"]  # what needs a look first
        assert client.get(f"/api/jobs/{job_id}/items", params={"sort": "id"}).status_code == 422


def test_previous_code_comes_back_step_by_step(isolated_api, monkeypatch):
    api = isolated_api
    monkeypatch.setattr(runtime.state, "catalog", FakeCatalog())
    _, [item_id] = seeded_job(api, [("needs_review", "OLD", 50.0)])
    with TestClient(api.app) as client:
        assert client.post(f"/api/items/{item_id}/undo-code").status_code == 400
        client.patch(f"/api/items/{item_id}", json={"ktru_code": "ONE"})
        client.patch(f"/api/items/{item_id}", json={"ktru_code": "TWO"})
        back = client.post(f"/api/items/{item_id}/undo-code").json()
        assert back["ktru_code"] == "ONE" and back["status"] == "corrected"
        first = client.post(f"/api/items/{item_id}/undo-code").json()
        assert first["ktru_code"] == "OLD" and first["status"] == "needs_review"
        assert client.post(f"/api/items/{item_id}/undo-code").status_code == 400


def test_accepted_rows_return_to_review_in_bulk(isolated_api):
    api = isolated_api
    job_id, ids = seeded_job(
        api,
        [
            ("accepted", "A", 90.0),
            ("accepted", "B", 70.0),
            ("auto", "C", 95.0),
            ("accepted", "D", 80.0),
        ],
    )
    with TestClient(api.app) as client:
        chosen = client.post(f"/api/jobs/{job_id}/revert", json={"ids": [ids[0], ids[2]]}).json()
        assert chosen == {"reverted": 1}  # the auto row is not accepted: nothing to return
        everything = client.post(f"/api/jobs/{job_id}/revert", json={"all_accepted": True}).json()
        assert everything == {"reverted": 2}
        counts = client.get(f"/api/jobs/{job_id}").json()["counts"]
        assert counts["accepted"] == 0
        assert client.post("/api/jobs/not-a-uuid/revert", json={"ids": [1]}).status_code == 404


def embeddings_gateway(monkeypatch, handler) -> None:
    client = httpx.Client(transport=httpx.MockTransport(handler))
    monkeypatch.setattr(embedder, "_client", client)
    monkeypatch.setattr(embedder, "YANDEX_KEY", "test-key")
    monkeypatch.setattr(embedder, "YANDEX_FOLDER", "b1gfolder")
    monkeypatch.setattr(embedder, "YANDEX_RPS", 1000.0)
    monkeypatch.setattr(embedder, "_cache", {})
    monkeypatch.setattr(embedder, "_status", {"up": None, "at": 0.0, "pause_until": 0.0})


def test_embeddings_state_follows_the_answers(monkeypatch):
    up = {"value": False}

    def handler(request):
        if not up["value"]:
            raise httpx.ConnectError("refused", request=request)
        return httpx.Response(200, json={"embedding": [1.0] * embedder.DIM})

    embeddings_gateway(monkeypatch, handler)
    assert embedder.state() == "down"  # the health probe asked the service itself
    assert embedder.embed(["Лампа"]) is None
    up["value"] = True
    assert embedder.embed(["Лампа"]).shape == (1, embedder.DIM)
    assert embedder.state() == "on"  # the job's answer is fresh: no second probe


def test_hung_embeddings_are_not_waited_for_again(monkeypatch):
    calls = []

    def handler(request):
        calls.append(request)
        raise httpx.ReadTimeout("no answer", request=request)

    embeddings_gateway(monkeypatch, handler)
    assert embedder.embed(["Лампа"]) is None
    assert embedder.embed(["Стол"]) is None  # paused: no second request
    assert len(calls) == 1 and embedder.state() == "down"


def test_broken_csv_is_a_clear_parse_error(tmp_path):
    path = tmp_path / "spec.csv"
    path.write_text("Наименование;Кол-во\n" + "Бумага " * 30_000 + ";1\n", encoding="utf-8")
    with pytest.raises(parser.ParseError, match="csv"):
        parser.parse_file(path)


class ActiveCodes:
    def get(self, code: str) -> dict:
        return {"code": code, "name": "Бумага", "status": "ACTIVE"}


def test_bulk_accept_skips_rows_that_are_not_products(isolated_api, monkeypatch):
    api = isolated_api
    monkeypatch.setattr(runtime.state, "catalog", ActiveCodes())
    job = models.Job(filename="spec.csv", status="done")
    with database.SessionLocal() as session:
        session.add(job)
        session.flush()
        for number, name in ((1, "Бумага"), (2, "")):  # row 2: «Итого по разделу» and the like
            item = models.Item(job_id=job.id, row_number=number, raw_text=name or "Раздел 1")
            item.attributes = {"name": name}
            item.ktru_code, item.confidence, item.status = "17.12.14.129-00000002", 99.0, "auto"
            session.add(item)
        session.commit()
    with TestClient(api.app) as client:
        assert client.post(f"/api/jobs/{job.id}/accept").json() == {"accepted": 1}


def test_cancelled_job_gets_no_failure_text(isolated_api):
    job = models.Job(filename="spec.csv", status="cancelled")
    with database.SessionLocal() as session:
        session.add(job)
        session.commit()
    job_store.finish(
        job.id, "failed", "Файл задания не найден: загрузите его заново", {"read": 0.1}
    )
    with database.SessionLocal() as session:
        stored = session.get(models.Job, job.id)
        assert stored.status == "cancelled" and stored.error is None
        assert stored.finished_at is not None


def test_item_page_says_how_candidates_differ(isolated_api):
    api = isolated_api

    def soap(code: str, mass: str) -> dict:
        return {
            "code": code,
            "name": "Мыло туалетное твердое",
            "score": 0.9,
            "characteristics": [["Вид", ["Туалетное"]], ["Масса нетто", [mass]]],
        }

    job = models.Job(filename="spec.csv", status="done")
    with database.SessionLocal() as session:
        session.add(job)
        session.flush()
        item = models.Item(job_id=job.id, row_number=2, raw_text="Мыло 100 г", status="auto")
        item.candidates = [soap("20.41.31.110-00000002", "< 90 Грамм"), soap("-07", "≥ 90 Грамм")]
        session.add(item)
        session.commit()
    with TestClient(api.app) as client:
        shown = client.get(f"/api/items/{item.id}").json()["candidates"]
    assert shown[0]["differs"] == [["Масса нетто", "< 90 Грамм"]]  # «Вид» is the same in both
    assert shown[1]["differs"] == [["Масса нетто", "≥ 90 Грамм"]]


def test_a_unit_the_customer_wrote_is_not_replaced_by_a_guess(isolated_api):
    guessed = {"measure": "усл. ед"}
    minute = parser.Item(1, "Услуги телефонной связи", {"Ед. изм.": "Минута"})
    assert pipeline.unit_hints(minute, guessed) == (None, None)  # a unit we do not know
    litre = parser.Item(2, "Кефир", {"Ед. изм.": "л"})
    assert pipeline.unit_hints(litre, {"measure": "кг"}) == ("l", "kg")
    blank = parser.Item(3, "Кефир", {"Ед. изм.": " "})
    assert pipeline.unit_hints(blank, {"measure": "л"}) == (None, "l")
    assert pipeline.unit_hints(parser.Item(4, "Кефир", {}), {"measure": "л"}) == (None, "l")


import database
import exporting
import job_store
import models
import routes
import routes.jobs
import runtime
import pipeline
