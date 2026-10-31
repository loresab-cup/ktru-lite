"""Odd input found by a run over every API address: control characters, numbers beyond the
database integer, unknown statuses. Each gets a plain answer instead of a server error."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import database as db
import exporting
import parser
from models import Item, Job


def make_job() -> tuple[str, int]:
    with db.SessionLocal() as session:
        job = Job(filename="spec.csv", status="done", total=1, processed=1)
        session.add(job)
        session.flush()
        item = Item(
            job_id=job.id,
            row_number=1,
            raw_text="Бумага А4",
            status="needs_review",
            ktru_code="17.12.14.129-00000002",
            ktru_name="Бумага",
            attributes={"name": "Бумага\x00 для\x07 печати"},  # stored before the fix
            candidates=[],
        )
        session.add(item)
        session.commit()
        return str(job.id), item.id


def test_control_characters_leave_file_cells_and_exported_cells():
    assert parser.clean("a\x00b\x07c\x1bd\x7fe  f") == "a b c d e f"
    assert exporting.cell("Бумага\x00\x0b А4") == "Бумага А4"
    assert exporting.cell("=1+1\x01") == "'=1+1"


def test_numbers_beyond_the_database_and_nul_in_searches_get_plain_answers(api_database):
    job_id, item_id = make_job()
    big = 2**70
    with TestClient(api_database.app) as client:
        assert client.get(f"/api/items/{big}").status_code == 404
        assert client.post(f"/api/items/{big}/accept").status_code == 404
        assert client.get(f"/api/items/{item_id}").status_code == 200
        revert = client.post(f"/api/jobs/{job_id}/revert", json={"ids": [big]})
        assert revert.status_code == 422
        items = client.get(f"/api/jobs/{job_id}/items", params={"offset": big})
        assert items.status_code == 422
        for path in ("/api/jobs", f"/api/jobs/{job_id}/items"):
            assert client.get(path, params={"q": "Бум\x00"}).status_code == 200


def test_unknown_statuses_are_ignored_in_the_row_filter(api_database):
    job_id, _ = make_job()
    with TestClient(api_database.app) as client:
        url = f"/api/jobs/{job_id}/items"
        assert client.get(url, params={"status": "needs_review,\x00"}).json()["total"] == 1
        assert client.get(url, params={"status": "nope"}).json()["total"] == 0


def test_export_survives_control_characters_already_in_the_database(api_database):
    job_id, _ = make_job()
    with TestClient(api_database.app) as client:
        answer = client.get(f"/api/jobs/{job_id}/export", params={"format": "xlsx"})
        assert answer.status_code == 200
        sheet = load_workbook(io.BytesIO(answer.content)).active
        texts = [str(c.value) for row in sheet.iter_rows() for c in row if c.value]
        assert not any(parser.CONTROL.search(text) for text in texts)
        assert client.get(f"/api/jobs/{job_id}/export", params={"format": "csv"}).status_code == 200


def test_typed_attributes_lose_control_characters(api_database, monkeypatch):
    import routes.items

    _, item_id = make_job()
    captured = {}

    class Catalog:
        def search(self, vectors, texts, include_inactive, attributes=None):
            captured["text"] = texts[0]
            captured["attributes"] = attributes
            return [[]]

    monkeypatch.setattr(routes.items, "require_catalog", lambda: Catalog())
    monkeypatch.setattr(routes.items.embedder, "embed", lambda texts: None)
    with TestClient(api_database.app) as client:
        body = {"name": "Бумага\x00", "characteristics": [{"name": "Формат\x07", "value": "А4"}]}
        answer = client.put(f"/api/items/{item_id}/attributes", json=body)
        assert answer.status_code == 200, answer.text
        assert answer.json()["attributes"]["name"] == "Бумага"
        assert "\x00" not in captured["text"] and "\x07" not in captured["text"]
        only_control = {"name": "\x00\x01", "characteristics": []}
        assert client.put(f"/api/items/{item_id}/attributes", json=only_control).status_code == 422
