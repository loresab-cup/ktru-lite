"""Scores the frozen run on all rows. Catalog names are read after the freeze, for difficulty statistics only."""
import collections,csv,gzip,json,random,re
from pathlib import Path
import openpyxl
O=Path('/out')
d=json.load(open(O/'run_raw.json',encoding='utf-8'));by={i['row_number']:i for i in d['items']}
ref=[json.loads(l) for l in open(O/'reference.jsonl',encoding='utf-8')]
ws=openpyxl.load_workbook(O/'reference_with_ktru.xlsx').active
H=[c.value for c in ws[1]];R=[dict(zip(H,r)) for r in ws.iter_rows(min_row=2,values_only=True)]
norm=lambda s:' '.join(re.findall(r'[0-9a-zа-я]+',(s or '').lower().replace('ё','е')))
cnt=collections.Counter();active=collections.Counter()
for l in gzip.open('/db/ktru_catalog.jsonl.gz','rt',encoding='utf-8'):
    c=json.loads(l);cnt[norm(c['name'])]+=1
    if 'ACTIVE' in str(c.get('status','')).upper():active[norm(c['name'])]+=1
rows=[]
for x,r in zip(ref,R):
    i=by.get(x['row_number'],{})
    rows.append(dict(n=x['row_number'],want=x['ktru_code'],got=i.get('ktru_code'),ok=i.get('ktru_code')==x['ktru_code'],status=i.get('status'),conf=i.get('confidence'),cat=r['Категория'],pur=r['Номер закупки'],name=x['name'],chars=r['Характеристики'],twins=active[norm(x['name'])],top5=x['ktru_code'] in [c['code'] for c in i.get('candidates') or []],gotname=i.get('ktru_name'),why=(i.get('attributes') or {}).get('review') or i.get('summary')))
N=len(rows);k=sum(r['ok'] for r in rows)
print('exact',k,N,round(100*k/N,1))
# bootstrap by purchase: rows of one purchase are not independent
pu=collections.defaultdict(list)
for r in rows:pu[r['pur']].append(r['ok'])
keys=list(pu);random.seed(1);v=[]
for _ in range(5000):
    a=[x for p in random.choices(keys,k=len(keys)) for x in pu[p]];v.append(100*sum(a)/len(a))
v.sort();print('95% CI by purchase',round(v[125],1),round(v[4875],1))
print('min share purchases',min(sum(x)/len(x) for x in pu.values()))
g=[('1 код с таким названием',lambda r:r['twins']<=1),('2-9 кодов',lambda r:1<r['twins']<10),('10+ кодов',lambda r:r['twins']>=10)]
for t,f in g:
    s=[r for r in rows if f(r)];print(t,len(s),sum(r['ok'] for r in s),round(100*sum(r['ok'] for r in s)/max(len(s),1),1))
for c,s in sorted(collections.Counter(r['cat'] for r in rows).items(),key=lambda x:-x[1]):
    print(c,s,sum(r['ok'] for r in rows if r['cat']==c))
print('purchases with errors',collections.Counter(r['pur'] for r in rows if not r['ok']).most_common())
print('auto',sum(r['status']=='auto' for r in rows),'auto wrong',sum(r['status']=='auto' and not r['ok'] for r in rows))
print('top5 miss',[r['n'] for r in rows if not r['top5']])
with open(O/'mismatches.csv','w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f,delimiter=';');w.writerow('строка закупка категория наименование характеристики код_заказчика код_сайта название_у_сайта статус уверенность в_топ5'.split())
    for r in rows:
        if not r['ok']:
            w.writerow([r['n'],r['pur'],r['cat'],r['name'],r['chars'],r['want'],r['got'],r['gotname'],r['status'],r['conf'],r['top5']])
            print('--',r['n'],r['pur'],r['name'][:40],'| want',r['want'],'| got',r['got'],r['status'],r['conf'],'|',r['chars'][:90])
# blind review sheets: no site answers inside
PEOPLE=['Артас','Евгений','Кирилл','Алексей','Егор']
pur=sorted(pu,key=lambda p:-len(pu[p]));load={p:[] for p in PEOPLE}
for p in pur:
    who=min(PEOPLE,key=lambda x:len(load[x]));load[who]+= [r for r in R if r['Номер закупки']==p]
Path(O/'review').mkdir(exist_ok=True)
for who,items in load.items():
    b=openpyxl.Workbook();s=b.active;s.title='Проверка'
    s.append(['№','Закупка','Ссылка','Наименование','Характеристики','Ед. изм.','Код заказчика','Название в каталоге','Код подходит? (да/нет/спорно)','Комментарий'])
    for r in items:s.append([r['№'],r['Номер закупки'],r['URL источника'],r['Наименование'],r['Характеристики'],r['Ед. изм.'],r['КТРУ'],r['Наименование КТРУ'],'',''])
    for col,w in zip('ABCDEFGHIJ',(6,22,30,34,70,8,24,34,16,30)):s.column_dimensions[col].width=w
    b.save(O/'review'/f'проверка_{who}.xlsx');print(who,len(items))
