"""Run real datasets against an isolated API; keep evidence without credentials.

python scripts/audit_live.py --api http://127.0.0.1:8001 --root dataset --output audit.json
The API must have its own database: this creates jobs and edits one of their rows.
"""

import argparse
import csv
import io
import json
import time
from pathlib import Path

import evaluate
import httpx


def check(response):
    response.raise_for_status()
    return response.json()


def dataset(client, path, reference, output):
    started = time.monotonic()
    job_id = evaluate.upload(client, path)
    job, items = evaluate.run(client, job_id)
    assert job["status"] == "done", (path.name, job["status"], job.get("error"))
    expected = {}
    if reference and reference.exists():
        expected = {
            row["row_number"]: row
            for row in map(json.loads, reference.read_text(encoding="utf-8").splitlines())
        }
        assert set(expected) <= {item["row_number"] for item in items}, path.name
    result = {
        "file": str(path),
        "job_id": job_id,
        "status": job["status"],
        "rows": len(items),
        "seconds": round(time.monotonic() - started, 2),
        "fallback": job.get("fallback"),
        "words_only": job.get("words_only"),
        "sources": {},
        "metrics": evaluate.evaluate(client, items, expected, False),
        "exports": {},
    }
    for item in items:
        attrs = item.get("attributes") or {}
        key = f"{attrs.get('source')}:{attrs.get('fallback', '')}"
        result["sources"][key] = result["sources"].get(key, 0) + 1
    for fmt in ("csv", "xlsx"):
        response = client.get(f"/api/jobs/{job_id}/export", params={"format": fmt})
        response.raise_for_status()
        target = output.parent / f"{path.stem}.{fmt}"
        target.write_bytes(response.content)
        if fmt == "csv":
            rows = list(
                csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";")
            )
            exported = len(rows) - 1
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(response.content), read_only=True)
            exported = workbook.active.max_row - 1
            workbook.close()
        assert exported == len(items), (path.name, fmt, exported, len(items))
        result["exports"][fmt] = exported
    (output.parent / f"{path.stem}.items.json").write_text(
        json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if items:
        item = items[0]
        code = item.get("ktru_code")
        if code:
            fixed = check(client.patch(f"/api/items/{item['id']}", json={"ktru_code": code}))
            assert fixed["confidence"] is None and fixed["status"] == "corrected"
            accepted = check(client.post(f"/api/items/{item['id']}/accept"))
            assert accepted["status"] == "accepted" and accepted["confidence"] is None
            restored = check(client.post(f"/api/items/{item['id']}/revert"))
            assert restored["ktru_code"] == code
            assert restored["confidence"] == item["confidence"]
            result["edit_accept_revert"] = True
    return result


def main():
    args = argparse.ArgumentParser(description=__doc__)
    args.add_argument("--api", default="http://127.0.0.1:8001")
    args.add_argument("--root", type=Path, default=Path("dataset"))
    args.add_argument("--output", type=Path, default=Path("audit-results/audit.json"))
    options = args.parse_args()
    options.output.parent.mkdir(parents=True, exist_ok=True)
    report = {"api": options.api, "datasets": []}
    files = [
        ("real_produkty", "produkty"),
        ("eis", "ООЗ_мыло_крем"),
        ("modified", "ООЗ_мыло_крем_изм"),
        ("organizers", "Спецификация_292"),
        ("modified", "Спецификация_292_изм"),
        ("organizers", "Спецификация_1000"),
        ("modified", "Спецификация_1000_изм"),
    ]
    with httpx.Client(base_url=options.api, timeout=120) as client:
        report["health_before"] = check(client.get("/api/health"))
        for folder, name in files:
            root = options.root / folder
            result = dataset(
                client,
                root / f"{name}_без_КТРУ.xlsx",
                root / f"{name}_reference.jsonl",
                options.output,
            )
            report["datasets"].append(result)
            options.output.write_text(
                json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            print(json.dumps(result, ensure_ascii=False), flush=True)
        synthetic = options.root / "synthetic" / "positions_1000.csv"
        report["datasets"].append(dataset(client, synthetic, None, options.output))
        report["health_after"] = check(client.get("/api/health"))
    options.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
