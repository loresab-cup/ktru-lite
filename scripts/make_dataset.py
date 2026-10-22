"""Reference files for scripts/evaluate.py, built from the «с_КТРУ» files of dataset/.

    python scripts/make_dataset.py [dataset]

Writes <name>_reference.jsonl next to every file with answers: one object per row
{"row_number", "ktru_code", "name", "unit"?, "characteristics": [{"name", "value"}]}.
"""

import json
import pathlib
import re
import sys

from openpyxl import load_workbook

CODE = re.compile(r"\d{2}\.\d{2}\.\d{2}\.\d{3}-\d{5,8}")


def sheet_rows(path: pathlib.Path) -> list[list[str]]:
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    return [
        ["" if v is None else str(v).strip() for v in row]
        for row in sheet.iter_rows(values_only=True)
    ]


def write(path: pathlib.Path, rows: list[dict]) -> None:
    text = "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n"
    path.write_text(text, encoding="utf-8")
    print(f"{path}: {len(rows)} rows")


def pairs(text: str) -> list[dict]:
    """«Диагональ экрана: ≥ 15.6 Дюйм\\nТип: SSD» -> [{"name", "value"}, ...]."""
    found = []
    for line in text.splitlines():
        name, colon, value = line.partition(":")
        if colon and name.strip() and value.strip():
            found.append({"name": name.strip(), "value": value.strip()})
    return found


def plain_reference(with_codes: pathlib.Path) -> list[dict]:
    """Organizers' format: A - name, B - «name: value» lines, C - code; no header."""
    reference = []
    for number, row in enumerate(sheet_rows(with_codes), start=1):
        code = "".join(row[2].split()) if len(row) > 2 else ""
        if not code:
            continue
        item = {"row_number": number, "ktru_code": code, "name": row[0], "characteristics": []}
        for pair in pairs(row[1]):
            if pair["name"].lower().startswith("единица измерения товара"):
                item["unit"] = pair["value"]
            else:
                item["characteristics"].append(pair)
        reference.append(item)
    return reference


def eis_unit(text: str) -> str:
    return text.replace("[3*]", "³").replace("[2*]", "²").split(";")[0].strip()


def column(header: list[str], *words: str) -> int:
    return next(i for i, t in enumerate(header) if all(w in t.lower() for w in words))


def eis_reference(with_codes: pathlib.Path) -> list[dict]:
    """EIS form: a product row with «Код: …» in the name cell, then one characteristic a row."""
    rows = sheet_rows(with_codes)
    start = next(
        i
        for i, row in enumerate(rows)
        if any("наименование характеристики" in c.lower() for c in row)
    )
    header = rows[start]
    name_at = column(header, "наименование товара")
    char_at = column(header, "наименование характеристики")
    values_at = [
        i for i, t in enumerate(header) if "значени" in t.lower() and "примен" not in t.lower()
    ]
    char_unit_at = column(header, "единица измерения характеристики")
    unit_at = column(header, "ед. изм")
    reference = []
    for number, row in enumerate(rows[start + 1 :], start=start + 2):
        code = CODE.search(row[name_at])
        if code:
            name = row[name_at].splitlines()[0].strip()
            reference.append(
                {"row_number": number, "ktru_code": code.group(), "name": name}
                | {"unit": eis_unit(row[unit_at]), "characteristics": []}
            )
        value = " ".join(row[i] for i in values_at if row[i])
        if reference and row[char_at] and value:
            unit = eis_unit(row[char_unit_at])
            reference[-1]["characteristics"].append(
                {"name": row[char_at].strip(), "value": f"{value} {unit}".strip()}
            )
    return reference


# --- modified copies: the same meaning written another way, as the judges' test files are ---

NUMBER = re.compile(r"\d+(?:[.,]\d+)?")
# Unit (lowercase) -> the same value in another unit: 100 Грамм -> 0,1 Килограмм
CONVERSIONS = {
    "грамм": ("Килограмм", 0.001),
    "килограмм": ("Грамм", 1000),
    "миллиметр": ("Сантиметр", 0.1),
    "сантиметр": ("Миллиметр", 10),
    "метр": ("Миллиметр", 1000),
    "литр": ("Миллилитр", 1000),
    "миллилитр": ("Литр", 0.001),
    "кубический сантиметр": ("Миллилитр", 1),
    "гигабайт": ("Мегабайт", 1024),
    "процент": ("%", 1),
}
WORD_OPS = {"≥": "не менее", "≤": "не более", ">": "более", "<": "менее"}


def number_text(value: float) -> str:
    return f"{value:.6f}".rstrip("0").rstrip(".").replace(".", ",")


def split_unit(value: str) -> tuple[str, str]:
    """«≥ 15.6 Дюйм (25,4 мм)» -> («≥ 15.6», «Дюйм (25,4 мм)»); no number -> (value, "")."""
    numbers = list(NUMBER.finditer(value.split("(")[0]))
    if not numbers:
        return value, ""
    end = numbers[-1].end()
    return value[:end].strip(), value[end:].strip()


def modify_value(value: str, n: int) -> str:
    """Even n: another unit; signs become words; every third plain number becomes a range."""
    body, unit = split_unit(value)
    if not unit and not NUMBER.search(body):
        return value.lower() if value in ("Да", "Нет") and n % 2 else value
    head, paren, tail = unit.partition("(")
    other = CONVERSIONS.get(head.strip().lower())
    if other and n % 2 == 0:
        head, factor = other
        body = NUMBER.sub(lambda m: number_text(float(m.group().replace(",", ".")) * factor), body)
    plain = not any(sign in body for sign in WORD_OPS)
    for sign, word in WORD_OPS.items():
        body = body.replace(sign, word)
    if plain and n % 3 == 0:
        body = f"не менее {body}"
    return " ".join(part for part in (body, head.strip(), paren + tail) if part)


def reword(name: str, n: int) -> str:
    """Every fourth name: words in another order, «Мыло туалетное» -> «Туалетное мыло»."""
    parts = name.split()
    # Prepositions («Крем для рук») and numbers do not survive a new word order
    if n % 4 != 1 or not 2 <= len(parts) <= 4 or min(len(p) for p in parts) < 4:
        return name
    reordered = " ".join(reversed(parts)).lower()
    return reordered[0].upper() + reordered[1:]


def save(rows: list[list], path: pathlib.Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    print(f"{path}: {len(rows)} rows")


def modify_plain(source: pathlib.Path, target: pathlib.Path) -> None:
    """Organizers' file with codes -> modified «с_КТРУ» and «без_КТРУ» copies and reference."""
    with_codes, without = [], []
    for n, row in enumerate(sheet_rows(source)):
        lines = []
        for k, line in enumerate(row[1].splitlines()):
            name, colon, value = line.partition(":")
            lines.append(f"{name}: {modify_value(value.strip(), n + k)}" if colon else line)
        name = reword(row[0], n)
        with_codes.append([name, "\n".join(lines), row[2]])
        without.append([name, "\n".join(lines)])
    stem = target.name
    save(with_codes, target.with_name(f"{stem}_с_КТРУ.xlsx"))
    save(without, target.with_name(f"{stem}_без_КТРУ.xlsx"))
    reference = plain_reference(target.with_name(f"{stem}_с_КТРУ.xlsx"))
    write(target.with_name(f"{stem}_reference.jsonl"), reference)


def modify_eis(source: pathlib.Path, target: pathlib.Path) -> None:
    """EIS form with codes -> the same form with other units, ranges and word order."""
    workbook = load_workbook(source)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows())
    start = next(
        i
        for i, row in enumerate(rows)
        if any("наименование характ" in str(c.value or "").lower() for c in row)
    )
    header = [str(c.value or "") for c in rows[start]]
    name_at = column(header, "наименование товара")
    char_at = column(header, "наименование характеристики")
    text_at = column(header, "текстовое описание")
    value_at = column(header, "конкретное значение")
    unit_at = column(header, "единица измерения характеристики")
    for n, row in enumerate(rows[start + 1 :]):
        name = str(row[name_at].value or "")
        if CODE.search(name):
            first, _, rest = name.partition("\n")
            row[name_at].value = f"{reword(first, 1)}\n{rest}"
        if not row[char_at].value:
            continue
        cell = row[value_at] if row[value_at].value is not None else row[text_at]
        unit = eis_unit(str(row[unit_at].value or ""))
        changed = modify_value(f"{cell.value} {unit}".strip(), n)
        cell.value, row[unit_at].value = split_unit(changed) if unit else (changed, None)
    with_codes = target.with_name(f"{target.name}_с_КТРУ.xlsx")
    workbook.save(with_codes)
    write(target.with_name(f"{target.name}_reference.jsonl"), eis_reference(with_codes))
    for row in rows[start:]:
        name = str(row[name_at].value or "")
        if name:  # empty ones may be parts of merged cells, which are read-only
            row[name_at].value = name.split("\n")[0].replace(", код КТРУ", "")
    workbook.save(target.with_name(f"{target.name}_без_КТРУ.xlsx"))


def main() -> None:
    root = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "dataset")
    organizers = root / "organizers"
    for n in ("1000", "292"):
        reference = plain_reference(organizers / f"Спецификация_{n}_с_КТРУ.xlsx")
        write(organizers / f"Спецификация_{n}_reference.jsonl", reference)
    real = root / "real_produkty"
    write(real / "produkty_reference.jsonl", plain_reference(real / "produkty_с_КТРУ.xlsx"))
    eis = root / "eis"
    write(eis / "ООЗ_мыло_крем_reference.jsonl", eis_reference(eis / "ООЗ_мыло_крем_с_КТРУ.xlsx"))

    modified = root / "modified"
    modified.mkdir(exist_ok=True)
    modify_eis(eis / "ООЗ_мыло_крем_с_КТРУ.xlsx", modified / "ООЗ_мыло_крем_изм")
    for n in ("1000", "292"):
        modify_plain(
            organizers / f"Спецификация_{n}_с_КТРУ.xlsx", modified / f"Спецификация_{n}_изм"
        )


if __name__ == "__main__":
    main()
