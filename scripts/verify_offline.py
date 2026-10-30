"""Verify a disposable stack with all external AI providers disabled; never use real accounts."""

import argparse
import csv
import io
import json
import secrets
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from docx import Document
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]


def checked(response, status=200):
    assert response.status_code == status, (response.status_code, response.text[:200])
    return response.json() if response.content else None


def upload(client, name, content):
    return checked(client.post("/api/jobs", files={"file": (name, content)}), 202)["job_id"]


def finished(client, job_id, timeout=180):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = checked(client.get(f"/api/jobs/{job_id}"))
        if job["status"] in ("done", "failed", "cancelled"):
            return job
        time.sleep(0.2)
    raise AssertionError("Job did not finish in time")


def items(client, job_id):
    result = []
    for offset in range(0, 1000, 200):
        page = checked(
            client.get(f"/api/jobs/{job_id}/items", params={"limit": 200, "offset": offset})
        )
        result.extend(page["items"])
        if len(result) >= page["total"]:
            break
    return result


def files():
    text = "Бумага А4\nЛампа 10 Вт"
    yield "sample.txt", text.encode("utf-16")
    yield "sample.csv", ("Наименование;Количество\nБумага А4;1\nЛампа 10 Вт;2").encode("utf-8-sig")
    book = Workbook()
    book.active.append(["Наименование", "Количество"])
    for name in text.splitlines():
        book.active.append([name, 1])
    data = io.BytesIO()
    book.save(data)
    book.close()
    yield "sample.xlsx", data.getvalue()
    document = Document()
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        [("Наименование", "Количество"), ("Бумага А4", "1"), ("Лампа 10 Вт", "2")],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    data = io.BytesIO()
    document.save(data)
    yield "sample.docx", data.getvalue()
    yield "sample.pdf", (ROOT / "backend/tests/data/table_spec.pdf").read_bytes()


def verify(api, output):
    report = {"api": api, "mode": "no external AI calls", "checks": {}, "risks": []}
    checks = report["checks"]
    with (
        httpx.Client(base_url=api, timeout=30) as client,
        httpx.Client(base_url=api, timeout=30) as other,
    ):
        health = checked(client.get("/api/health"))
        assert health["catalog"] == "ready"
        assert health["llm"] == health["embeddings"] == "off", "Refusing a potentially paid run"
        assert health["codes"] == 81509 and health["vector_search"] == "pgvector"
        checks["health"] = {
            key: health[key]
            for key in ("catalog", "codes", "llm", "embeddings", "queue", "vector_search")
        }
        policy = checked(client.get("/api/legal"))
        consent = {"privacy_version": policy["version"], "privacy_accepted": True}
        checked(client.post("/api/auth/guest", json={}), 400)
        response = client.post("/api/auth/guest", json=consent)
        assert checked(response, 201)["is_guest"]
        cookie = response.headers["set-cookie"].lower()
        assert "httponly" in cookie and "samesite=lax" in cookie and "max-age" not in cookie
        checks["guest_and_consent"] = True

        jobs = []
        for name, content in files():
            job_id = upload(client, name, content)
            job = finished(client, job_id)
            assert job["status"] == "done", (name, job.get("error"))
            rows = items(client, job_id)
            assert len(rows) == job["total"] >= 2
            assert job["words_only"] == job["total"]
            checks[name] = {"rows": len(rows), "seconds": job["stage_seconds"]["total"]}
            jobs.append((job_id, rows))

        job_id, rows = jobs[0]
        row = rows[0]
        assert len(row["candidates"]) >= 2
        replacement = next(c["code"] for c in row["candidates"] if c["code"] != row["ktru_code"])
        saved = checked(client.patch(f"/api/items/{row['id']}", json={"ktru_code": replacement}))
        assert saved["ktru_code"] == replacement and saved["confidence"] is None
        assert (
            checked(client.post(f"/api/items/{row['id']}/undo-code"))["ktru_code"]
            == row["ktru_code"]
        )
        checked(client.patch(f"/api/items/{row['id']}", json={"ktru_code": replacement}))
        checked(client.post(f"/api/items/{rows[1]['id']}/accept"))
        for fmt in ("csv", "xlsx"):
            response = client.get(f"/api/jobs/{job_id}/export", params={"format": fmt})
            assert response.status_code == 200
            if fmt == "csv":
                exported = list(
                    csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";")
                )
            else:
                workbook = load_workbook(io.BytesIO(response.content), read_only=True)
                exported = list(workbook.active.values)
                workbook.close()
            assert len(exported) == len(rows) + 1
            assert replacement in exported[1]
        checks["manual_edit_undo_accept_exports"] = True

        password = secrets.token_urlsafe(20)
        login = "verify-" + secrets.token_hex(5)
        checked(
            client.post(
                "/api/auth/register", json=consent | {"login": login, "password": password}
            ),
            201,
        )
        assert checked(client.get(f"/api/jobs/{job_id}"))["id"] == job_id
        repeated = finished(client, upload(client, "remember.txt", row["raw_text"].encode()))
        remembered = items(client, repeated["id"])[0]
        assert remembered["ktru_code"] == replacement and remembered["attributes"].get("remembered")
        checks["guest_adoption_and_memory"] = True
        checked(
            other.post(
                "/api/auth/register",
                json=consent
                | {
                    "login": "verify-" + secrets.token_hex(5),
                    "password": secrets.token_urlsafe(20),
                },
            ),
            201,
        )
        for path in (
            f"/api/jobs/{job_id}",
            f"/api/jobs/{job_id}/items",
            f"/api/jobs/{job_id}/export",
            f"/api/items/{row['id']}",
        ):
            checked(other.get(path), 404)
        checked(other.patch(f"/api/items/{row['id']}", json={"ktru_code": replacement}), 404)
        checks["account_isolation"] = True
        checked(client.post("/api/auth/logout"), 204)
        checked(client.get("/api/jobs"), 401)
        checked(
            client.post("/api/auth/login", json=consent | {"login": login, "password": password})
        )
        checks["login_logout"] = True

        for name, content, expected in [
            ("empty.txt", b"", 400),
            ("wrong.pdf", b"text", 400),
            ("large.txt", b"x" * (10 * 1024 * 1024 + 1), 413),
        ]:
            checked(client.post("/api/jobs", files={"file": (name, content)}), expected)
        oversized = finished(client, upload(client, "1001.txt", ("Бумага А4\n" * 1001).encode()))
        assert oversized["status"] == "failed" and "1000" in oversized["error"]
        checks["invalid_empty_oversize_row_limit"] = True

        with ThreadPoolExecutor(2) as pool:
            pending = list(
                pool.map(lambda n: upload(client, f"parallel-{n}.txt", b"Paper A4"), range(2))
            )
        assert all(finished(client, j)["status"] == "done" for j in pending)
        checks["parallel_queue"] = True

        probe = finished(client, upload(client, "formula.txt", b"=1+1"))
        formula = items(client, probe["id"])[0]
        if formula["ktru_code"] is not None:
            report["risks"].append(
                "Formula-only input still receives a catalog code; classification safeguards were not merged without full A/B inputs."
            )
        auto = sum(row["status"] == "auto" for _, rows in jobs for row in rows)
        if auto:
            report["risks"].append(
                f"Word-only fallback assigned auto status to {auto} rows of the format probes."
            )
        checks["formula_export_escaped"] = (
            client.get(f"/api/jobs/{probe['id']}/export", params={"format": "csv"})
            .content.decode("utf-8-sig")
            .find("'=1+1")
            >= 0
        )
        assert checks["formula_export_escaped"]
        report["sample_job"] = job_id
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--api", default="http://127.0.0.1:18000")
    parser.add_argument("--out", type=Path, default=ROOT / "outputs/final-prep/offline-api.json")
    args = parser.parse_args()
    target = urlsplit(args.api)
    if target.hostname not in ("127.0.0.1", "localhost") or target.scheme != "http":
        parser.error("Use a disposable stack on localhost")
    verify(args.api, args.out)


if __name__ == "__main__":
    main()
